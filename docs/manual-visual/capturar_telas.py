"""Capture the real macOS UI using temporary, fictional DOCX fixtures only.

Run from the repository root; see README.md for documentation-only dependencies.
No changes to production configuration or source documents are made.
"""
from pathlib import Path
from tempfile import TemporaryDirectory
import json
import time
from tkinter import ttk

import AppKit
import pymupdf as fitz
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.ui import FormatWordApp

HERE = Path(__file__).resolve().parent
OUT = HERE / 'telas'
OUT.mkdir(exist_ok=True)
metadata = {}


def pump(app, seconds=.4):
    end = time.monotonic() + seconds
    while time.monotonic() < end:
        app.update()
        time.sleep(.02)


def wait_for(app, predicate):
    end = time.monotonic() + 30
    while not predicate():
        assert time.monotonic() < end, 'UI operation timed out'
        pump(app, .1)
    pump(app)


def descendants(widget):
    for child in widget.winfo_children():
        yield child
        yield from descendants(child)


def capture(app, name, window=None):
    window = window or app
    # Force a full native repaint after tab switches; otherwise Tk's cached
    # backing view can still contain fragments of the previous tab.
    window.withdraw()
    pump(app, .15)
    window.deiconify()
    window.lift()
    pump(app, .8)
    native = next(w for w in AppKit.NSApplication.sharedApplication().windows()
                  if str(w.title()) == window.title())
    view = native.contentView()
    rect = view.bounds()
    bitmap = view.bitmapImageRepForCachingDisplayInRect_(rect)
    view.cacheDisplayInRect_toBitmapImageRep_(rect, bitmap)
    data = bitmap.representationUsingType_properties_(AppKit.NSBitmapImageFileTypePNG, {})
    assert data.writeToFile_atomically_(str(OUT / (name + '.png')), True)
    widgets = []
    for widget in descendants(window):
        if not widget.winfo_ismapped():
            continue
        try:
            label = str(widget.cget('text'))
        except Exception:
            continue
        if label:
            widgets.append(dict(text=label, x=widget.winfo_rootx()-window.winfo_rootx(),
                y=widget.winfo_rooty()-window.winfo_rooty(),
                w=widget.winfo_width(), h=widget.winfo_height()))
    metadata[name] = dict(width=window.winfo_width(), height=window.winfo_height(),
                          pixels=[bitmap.pixelsWide(), bitmap.pixelsHigh()], widgets=widgets)
    print('Captured', name, flush=True)


def make_doc(path, picture=None, table=False, reference=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.5)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    title = doc.styles['Heading 1']
    title.font.name, title.font.size, title.font.bold = 'Arial', Pt(16), True
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    quote = doc.styles['Quote']
    quote.font.name, quote.font.size = 'Arial', Pt(10)
    quote.paragraph_format.left_indent = Cm(2)
    quote.paragraph_format.first_line_indent = Cm(0)
    doc.add_paragraph('Relatório de atividades' if reference else 'Solicitação de materiais',
                      style='Heading 1' if reference else 'Normal')
    doc.add_paragraph('Este documento fictício demonstra o uso do Format Word. '
                      'Os nomes e valores são apenas exemplos para o manual.')
    doc.add_paragraph('A equipe solicita a organização dos materiais para a próxima reunião.')
    if reference:
        doc.add_paragraph('Trecho de exemplo com formatação de citação.', style='Quote')
    if picture:
        doc.add_picture(str(picture), width=Cm(7))
    if table:
        cells = doc.add_table(rows=2, cols=2).rows
        cells[0].cells[0].text, cells[0].cells[1].text = 'Item', 'Quantidade'
        cells[1].cells[0].text, cells[1].cells[1].text = 'Caderno', '10'
    doc.add_paragraph('Equipe de exemplo')
    doc.save(path)


def main():
    with TemporaryDirectory(prefix='FW-', dir='/tmp') as tmp:
        root = Path(tmp)
        inputs = root / 'Documentos'
        inputs.mkdir()
        # Native vector demo figure, rendered solely as a synthetic DOCX fixture.
        figure = fitz.open()
        page = figure.new_page(width=360, height=120)
        page.draw_rect(fitz.Rect(0, 0, 360, 120), fill=(.93, .96, 1), color=None)
        for i, height in enumerate((25, 48, 72)):
            page.draw_rect(fitz.Rect(50+i*90, 100-height, 95+i*90, 100),
                           fill=(.13, .34, .70), color=None)
        page.get_pixmap(matrix=fitz.Matrix(2, 2)).save(root / 'grafico.png')
        figure.close()
        reference = root / 'Modelo-do-escritorio.docx'
        make_doc(reference, reference=True)
        first = inputs / 'Solicitacao.docx'
        second = inputs / 'Documento-com-tabela.docx'
        invalid = inputs / 'Arquivo-invalido.docx'
        make_doc(first, picture=root / 'grafico.png')
        make_doc(second, table=True)
        invalid.write_text('Exemplo fictício de arquivo inválido; não é um pacote DOCX.')

        app = FormatWordApp(config_dir=root / 'config')
        try:
            app.geometry('1000x760+40+40')
            pump(app)
            capture(app, '01-inicio')
            app.tabs.set('Perfis e formatação')
            capture(app, '02-perfis')
            assert app.import_profile(reference)
            wait_for(app, lambda: app.import_dialog is not None and not app.busy)
            dialog = app.import_dialog
            capture(app, '03-importar', dialog)
            notebook = next(w for w in descendants(dialog) if isinstance(w, ttk.Notebook))
            notebook.select(1)
            capture(app, '04-avisos-importacao', dialog)
            dialog.confirm()
            app.profile_name.set('Escritório — documentos gerais')
            # Actual edits through the form: round reference measurements to
            # the clear example values described in the guide.
            for name, value in {'first_line_indent_cm': '1,25', 'margin_top_cm': '2,5',
                                'margin_bottom_cm': '2,5', 'margin_left_cm': '2,5',
                                'margin_right_cm': '2,5'}.items():
                app.field_vars[name].set(value)
            capture(app, '05-salvar-perfil')
            assert app.save_profile()

            app.geometry('1000x950+40+30')
            for name, section in [('06-fonte', 'Fonte'), ('07-paragrafo', 'Parágrafo'),
                                  ('08-pagina', 'Página'), ('09-cabecalho', 'Cabeçalho e rodapé'),
                                  ('10-imagens', 'Imagens'), ('11-avancadas', 'Opções avançadas')]:
                app.show_profile_section(section)
                capture(app, name)
                detail_field = {'Fonte': 'font_color', 'Parágrafo': 'right_indent_cm',
                                'Página': 'margin_right_cm', 'Cabeçalho e rodapé': 'footer_distance_cm',
                                'Opções avançadas': 'max_input_mb'}.get(section)
                if detail_field:
                    app._scroll_to_field(detail_field)
                    if section in {'Fonte', 'Opções avançadas'}:
                        app.form._parent_canvas.yview_moveto(1)
                    capture(app, name + '-detalhe')
            app.geometry('1000x760+40+40')
            app.show_profile_section('Fonte')
            rules = app.edit_category_rules()
            rules.category_var.set('Título')
            rules.switch_category()
            rules.mode_var.set('Personalizar')
            rules.refresh_controls()
            capture(app, '12-tipos', rules)
            rules.confirm()
            assert app.save_profile()

            app.tabs.set('Arquivos e resultados')
            app.add_files([first, second, invalid])
            app.output_dir.set(str(root / 'Formatados'))
            capture(app, '13-arquivos')
            app.file_tree.selection_set('0')
            review = app.review_document()
            assert review is not None
            review.paragraph_tree.selection_set('0')
            review.category_var.set('Título')
            capture(app, '14-revisar-textos', review)
            review.classify_selected()
            notebook = next(w for w in descendants(review) if isinstance(w, ttk.Notebook))
            notebook.select(1)
            review.image_tree.selection_set(review.image_tree.get_children()[0])
            review.select_image()
            review.width_var.set('8')
            review.alignment_var.set('Centralizado')
            capture(app, '15-revisar-imagens', review)
            review.confirm()
            capture(app, '16-pronto-aplicar')
            assert app.start_batch()
            wait_for(app, lambda: not app.busy)
            statuses = [app.file_tree.set(str(i), 'status') for i in range(3)]
            assert statuses == ['Concluído', 'Concluído com ressalvas', 'Erro'], statuses
            app.file_tree.selection_set('1')
            app._show_result()
            capture(app, '17-resultados')
            app.file_tree.selection_set('2')
            app._show_result()
            capture(app, '18-erro')
            metadata['validation'] = dict(statuses=statuses, fictional_data=True,
                                           original_files_intact=all(p.exists() for p in [first, second, invalid]))
        finally:
            app.destroy()
    (HERE / 'capturas.json').write_text(json.dumps(metadata, ensure_ascii=False, indent=2)+'\n')


if __name__ == '__main__':
    main()
