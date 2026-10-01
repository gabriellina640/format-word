"""Small executable smoke test using only temporary documents and profiles."""
from __future__ import annotations

import json
import tempfile
import time
import traceback
from pathlib import Path


def run_self_test(report_path: Path, *, gui: bool = True) -> int:
    report = {'ok': False, 'checks': []}
    try:
        from docx import Document
        from PIL import Image
        from app.batch import format_documents_batch
        from app.config import FormatSettings
        with tempfile.TemporaryDirectory(prefix='formatword-selftest-') as directory:
            root = Path(directory)
            picture = root / 'picture.png'
            Image.new('RGB', (200, 100), 'navy').save(picture)
            source = root / 'sample.docx'
            document = Document()
            document.add_heading('Título de teste', 1)
            document.add_paragraph('Corpo do documento')
            document.add_picture(str(picture))
            document.save(source)
            original = source.read_bytes()
            settings = FormatSettings(formatting_mode='by_category', body_image_width_cm=4,
                category_rules={'title': {'mode':'custom','values':{'font_size':18,'alignment':'center'}}})
            item = format_documents_batch([source],root/'output',settings)[0]
            if item.error or item.result is None:
                raise RuntimeError(item.error or 'Saída ausente')
            output = Document(item.result.output_path)
            if (output.paragraphs[0].runs[0].font.size.pt != 18
                    or output.paragraphs[1].runs[0].font.size.pt != 12
                    or abs(output.inline_shapes[0].width.cm - 4) > .001
                    or abs(output.inline_shapes[0].height.cm - 2) > .001
                    or source.read_bytes() != original):
                raise RuntimeError('Saída diferente das propriedades solicitadas')
            report['checks'].append('DOCX, categorias, imagem proporcional, lote e original intacto')
            from app.profile_import import inspect_profile
            imported = inspect_profile(item.result.output_path).build_settings()
            if (imported.font_size != 12
                    or imported.category_rules['title']['values']['font_size'] != 18):
                raise RuntimeError('Importação não recuperou as propriedades da referência')
            reapplied = format_documents_batch([source], root / 'imported-output', imported)[0]
            if reapplied.error or reapplied.result is None:
                raise RuntimeError(reapplied.error or 'Saída do perfil importado ausente')
            check = Document(reapplied.result.output_path)
            if check.paragraphs[0].runs[0].font.size.pt != 18 or check.paragraphs[1].runs[0].font.size.pt != 12:
                raise RuntimeError('Perfil importado não reproduziu fonte por categoria')
            report['checks'].append('Importação de perfil DOCX e aplicação em outro documento')
            if gui:
                from app.ui import FormatWordApp
                app = FormatWordApp(config_dir=root/'profiles')
                try:
                    app.withdraw()
                    app.update()
                    app.tabs.set('Perfis e formatação')
                    app.show_profile_section('Parágrafo')
                    app.spacing_buttons[1].invoke()
                    indent = app.field_widgets['first_line_indent_cm']
                    indent.kind.set('Deslocado')
                    indent.selector.event_generate('<<ComboboxSelected>>')
                    indent.amount.set('0,75')
                    if app._parse().first_line_indent_cm != -.75:
                        raise RuntimeError('Controle de recuo especial não converteu Deslocado')
                    app.show_profile_section('Fonte')
                    app.add_files([source])
                    app.profile_name.set('Perfil de teste')
                    if not app.save_profile():
                        raise RuntimeError('Falha ao salvar perfil pela interface')
                    app.update()
                    app.field_widgets['formatting_mode']._dropdown_callback('Por tipo de texto')
                    if app.field_vars['formatting_mode'].get() != 'Por tipo de texto':
                        raise RuntimeError('Seletor não aceitou a primeira seleção')
                    app.field_widgets['formatting_mode']._dropdown_callback('Todo o texto')
                    if app.field_vars['formatting_mode'].get() != 'Todo o texto':
                        raise RuntimeError('Seletor não aceitou a troca de volta')
                    if not app.save_profile() or not app.import_profile(item.result.output_path):
                        raise RuntimeError('Não foi possível abrir a importação pela interface')
                    deadline = time.monotonic() + 20
                    while app.import_dialog is None and app.busy and time.monotonic() < deadline:
                        app.update()
                        time.sleep(.01)
                    if app.import_dialog is None:
                        raise RuntimeError('Revisão da importação não abriu: ' + app.status_text.get())
                    app.import_dialog.confirm()
                    if not app.dirty or not app.save_profile():
                        raise RuntimeError('Perfil importado não pôde ser salvo pela interface')
                finally:
                    app.destroy()
                report['checks'].append('Interface Tk, navegação guiada, recuo especial, seletores, revisão da importação e persistência isolada')
        report['ok'] = True
    except Exception:
        report['error'] = traceback.format_exc()
    report_path = Path(report_path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return 0 if report['ok'] else 1
