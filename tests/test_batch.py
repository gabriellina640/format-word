from __future__ import annotations

import tempfile
import threading
import unittest
from pathlib import Path
from docx import Document
from app.config import FormatSettings
import app.batch as batch


class BatchTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.paths = []
        for index in range(3):
            source = self.root / f'{index}.docx'
            doc = Document()
            doc.add_paragraph(f'Arquivo {index}')
            doc.save(source)
            self.paths.append(source)

    def test_continues_after_invalid_file_and_reports_each_result(self):
        self.paths[1].write_text('Inválido')
        events = []
        results = batch.format_documents_batch(self.paths,self.root/'out',FormatSettings(),on_result=events.append)
        self.assertEqual(len(results),3)
        self.assertEqual(results,events)
        self.assertIsNotNone(results[0].result)
        self.assertIsNotNone(results[1].error)
        self.assertIsNotNone(results[2].result)
        self.assertEqual(Document(results[2].result.output_path).paragraphs[0].text,'Arquivo 2')

    def test_cancellation_between_files_keeps_completed_output(self):
        cancel = threading.Event()
        results = batch.format_documents_batch(self.paths,self.root/'out',FormatSettings(),
            cancel_event=cancel,on_result=lambda _:cancel.set())
        self.assertIsNotNone(results[0].result)
        self.assertTrue(results[1].cancelled)
        self.assertTrue(results[2].cancelled)
        self.assertEqual(len(list((self.root/'out').glob('*.docx'))),1)

    def test_settings_snapshot_is_not_changed_by_callbacks(self):
        settings = FormatSettings(font_size=12.5)
        results = batch.format_documents_batch(self.paths,self.root/'out',settings,
            on_result=lambda _:setattr(settings,'font_size',36))
        for item in results:
            self.assertEqual(Document(item.result.output_path).paragraphs[0].runs[0].font.size.pt,12.5)

    def test_duplicate_filenames_from_distinct_folders_have_unique_outputs(self):
        other = self.root/'other'
        other.mkdir()
        copy = other/self.paths[0].name
        copy.write_bytes(self.paths[0].read_bytes())
        results = batch.format_documents_batch([self.paths[0],copy],self.root/'out',FormatSettings())
        self.assertNotEqual(results[0].result.output_path,results[1].result.output_path)

    def test_canceled_before_start_does_not_create_outputs(self):
        cancel = threading.Event()
        cancel.set()
        results = batch.format_documents_batch(self.paths,self.root/'out',FormatSettings(),cancel_event=cancel)
        self.assertTrue(all(r.cancelled for r in results))
        self.assertFalse((self.root/'out').exists())

    def test_empty_batch_is_rejected(self):
        with self.assertRaisesRegex(ValueError,'arquivo'):
            batch.format_documents_batch([],self.root/'out',FormatSettings())

    def test_invalid_global_settings_does_not_publish_any_file(self):
        with self.assertRaises(ValueError):
            batch.format_documents_batch(self.paths,self.root/'out',FormatSettings(font_size=0))
        self.assertFalse((self.root/'out').exists())

    def test_review_by_file_and_nested_snapshot(self):
        from app.document_model import DocumentOverrides, inspect_document
        settings = FormatSettings(formatting_mode='by_category', category_rules={
            'signature': {'mode': 'custom', 'values': {'font_size': 18}}})
        edits = {path: DocumentOverrides(inspect_document(path, settings).fingerprint,
                 {0: 'signature'}) for path in self.paths}
        def mutate(_):
            settings.category_rules['signature']['values']['font_size'] = 30
            for edit in edits.values():
                edit.paragraph_categories.clear()
        results = batch.format_documents_batch(self.paths, self.root/'out', settings,
            on_result=mutate, overrides_by_path=edits)
        self.assertTrue(all(item.error is None for item in results))
        for item in results:
            self.assertEqual(Document(item.result.output_path).paragraphs[0].runs[0].font.size.pt, 18)

    def test_stale_review_fails_only_its_file(self):
        from app.document_model import DocumentOverrides, inspect_document
        edits = {self.paths[1]: DocumentOverrides(inspect_document(self.paths[1], FormatSettings()).fingerprint)}
        doc = Document(self.paths[1]); doc.add_paragraph('Novo'); doc.save(self.paths[1])
        results = batch.format_documents_batch(self.paths, self.root/'out', FormatSettings(), overrides_by_path=edits)
        self.assertIsNotNone(results[0].result)
        self.assertIn('mudou', results[1].error)
        self.assertIsNotNone(results[2].result)


if __name__ == '__main__':
    unittest.main()
