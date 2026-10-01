"""Real Tk integration; run with FORMATWORD_GUI_TESTS=1 on a desktop session."""
import os
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch

from docx import Document


@unittest.skipUnless(os.environ.get('FORMATWORD_GUI_TESTS') == '1', 'Requires desktop Tk session')
class UIIntegrationTests(unittest.TestCase):
    def setUp(self):
        from app.ui import FormatWordApp
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()
        self.app = FormatWordApp(config_dir=self.root / 'config')
        self.app.update()

    def tearDown(self):
        if not self.app._destroyed:
            self.app.destroy()
        self.temp.cleanup()

    def pump_until(self, predicate, timeout=15):
        deadline = time.monotonic() + timeout
        while not predicate() and time.monotonic() < deadline:
            self.app.update()
            time.sleep(.01)
        self.assertTrue(predicate(), 'Tk operation exceeded timeout')

    def source(self, name='source.docx'):
        path = self.root / name
        doc = Document()
        doc.add_paragraph('Conteúdo original')
        doc.save(path)
        return path

    def test_dropdown_choices_take_effect_on_first_selection(self):
        from app.ui_fields import FIELDS
        # Exercise the actual menu path: CTk deletes and inserts in its Entry.
        # StringVar.set alone misses the re-entrant readonly-state regression.
        self.app.field_vars['header_mode'].set('Imagem do perfil')
        self.app.field_vars['footer_mode'].set('Imagem do perfil')
        self.app.update()
        for field in FIELDS:
            if not field.choices:
                continue
            widget = self.app.field_widgets[field.name]
            if field.name in ('header_alignment', 'footer_alignment'):
                self.app.field_vars[field.name.split('_')[0] + '_mode'].set('Imagem do perfil')
                self.app.update()
            for value in list(field.choices) + list(reversed(field.choices)):
                with self.subTest(field=field.name, value=value):
                    widget._dropdown_callback(value)
                    self.assertEqual(self.app.field_vars[field.name].get(), value)
                    self.app.update()
                    self.assertEqual(widget.get(), value)

    def test_dropdown_dependent_controls_busy_and_pending_close(self):
        self.app.field_widgets['header_mode']._dropdown_callback('Imagem do perfil')
        self.app.update()
        self.assertEqual(self.app.field_widgets['header_alignment'].cget('state'), 'readonly')
        self.app.field_widgets['header_mode']._dropdown_callback('Preservar original')
        self.app.update()
        self.assertEqual(self.app.field_widgets['header_alignment'].cget('state'), 'disabled')
        self.app.field_vars['alignment'].set('Centralizado')
        self.app._set_busy(True)
        self.app.update()
        self.assertEqual(self.app.field_widgets['alignment'].cget('state'), 'disabled')
        self.app._set_busy(False)
        self.app.field_vars['alignment'].set('Direita')
        self.app.destroy()
        self.assertFalse(self.app._after_ids)

    def test_guided_sections_keep_values_and_reveal_invalid_field(self):
        self.app.tabs.set('Perfis e formatação')
        self.assertEqual(self.app.profile_section.get(), 'Fonte')
        self.app.field_vars['font_size'].set('14,5')
        self.app.show_profile_section('Página')
        self.app.update()
        self.assertTrue(self.app.field_widgets['margin_left_cm'].winfo_ismapped())
        self.assertFalse(self.app.field_widgets['font_size'].winfo_ismapped())
        self.app.move_profile_section(-1)
        self.assertEqual(self.app.profile_section.get(), 'Parágrafo')
        self.assertEqual(self.app.field_vars['font_size'].get(), '14,5')
        self.app.field_vars['max_input_mb'].set('abc')
        self.assertFalse(self.app.save_profile())
        self.assertEqual(self.app.profile_section.get(), 'Opções avançadas')
        self.app.update()
        self.assertTrue(self.app.field_widgets['max_input_mb'].winfo_ismapped())

    def test_word_spacing_presets_indent_and_color_cancel(self):
        self.app.show_profile_section('Parágrafo')
        for button, expected in zip(self.app.spacing_buttons, (1, 1.5, 2)):
            button.invoke()
            self.app.update()
            self.assertEqual(self.app._parse().line_spacing, expected)
            self.assertEqual(self.app._parse().line_spacing_mode, 'multiple')
        control = self.app.field_widgets['first_line_indent_cm']
        control.kind.set('Deslocado')
        control.selector.event_generate('<<ComboboxSelected>>')
        control.amount.set('0,75')
        self.assertEqual(self.app._parse().first_line_indent_cm, -.75)
        self.app.field_vars['font_color'].set('123456')
        with patch('app.ui.colorchooser.askcolor', return_value=(None, None)):
            self.app.choose_font_color()
        self.assertEqual(self.app.field_vars['font_color'].get(), '123456')
        with patch('app.ui.colorchooser.askcolor', return_value=((17, 34, 51), '#112233')):
            self.app.choose_font_color()
        self.assertEqual(self.app._parse().font_color, '112233')
        self.app._set_busy(True)
        self.assertTrue(all(str(button.cget('state')) == 'disabled' for button in self.app.spacing_buttons))
        self.app._set_busy(False)

    def test_category_editor_uses_word_controls_without_losing_precision(self):
        self.app.new_profile()
        dialog = self.app.edit_category_rules()
        dialog.category_var.set('Citação')
        dialog.switch_category()
        dialog.mode_var.set('Personalizar')
        dialog.refresh_controls()
        control = dialog.controls['first_line_indent_cm']
        control.kind.set('Deslocado')
        control.selector.event_generate('<<ComboboxSelected>>')
        control.amount.set('0,7654321')
        dialog.spacing_buttons[2].invoke()
        with patch('app.review_dialogs.colorchooser.askcolor', return_value=((17, 34, 51), '#112233')):
            dialog.choose_font_color()
        dialog.confirm()
        self.assertFalse(dialog.winfo_exists())
        values = self.app._draft_base.category_rules['quote']['values']
        self.assertEqual(values['first_line_indent_cm'], -.7654321)
        self.assertEqual(values['line_spacing_mode'], 'multiple')
        self.assertEqual(values['line_spacing'], 2)
        self.assertEqual(values['font_color'], '112233')
        self.assertTrue(self.app.save_profile())
        self.assertEqual(self.app.store.load().stacks[self.app._selected].category_rules['quote']['values'], values)

    def test_profile_navigation_fits_minimum_window_and_keyboard(self):
        from app.ui import PROFILE_SECTIONS
        self.app.geometry('760x600')
        self.app.tabs.set('Perfis e formatação')
        self.app.update()
        for section in PROFILE_SECTIONS:
            self.app.show_profile_section(section)
            self.app.update()
            self.assertGreaterEqual(self.app.form._parent_canvas.winfo_height(), 70)
            for widget in (self.app.section_selector, self.app.previous_section_button, self.app.next_section_button):
                self.assertLessEqual(widget.winfo_rootx() - self.app.winfo_rootx() + widget.winfo_width(), self.app.winfo_width())
                self.assertLessEqual(widget.winfo_rooty() - self.app.winfo_rooty() + widget.winfo_height(), self.app.winfo_height())
        button = self.app.section_selector._buttons_dict['Parágrafo']
        self.app.focus_force()
        button._canvas.focus_force()
        self.app.update()
        button._canvas.event_generate('<Return>')
        self.app.update()
        self.assertEqual(self.app.profile_section.get(), 'Parágrafo')

    def test_word_shortcut_buttons_scroll_fully_into_view_on_focus(self):
        self.app.geometry('760x600')
        self.app.tabs.set('Perfis e formatação')
        self.app.focus_force()
        self.app.update()
        from app.ui import field_section
        from app.ui_fields import FIELD_MAP
        for name in ('font_color', 'line_spacing_mode'):
            self.app.show_profile_section(field_section(FIELD_MAP[name]))
            self.app._scroll_to_field(name)
            buttons = self.app.color_buttons if name == 'font_color' else self.app.spacing_buttons
            for button in buttons:
                target = getattr(button, '_canvas', button)
                target.focus_force()
                self.app.update()
                canvas = self.app.form._parent_canvas
                top = button.winfo_rooty() - canvas.winfo_rooty()
                self.assertGreaterEqual(top, 0)
                self.assertLessEqual(top + button.winfo_height(), canvas.winfo_height())

    def test_import_cancel_preserves_unsaved_profile(self):
        self.app.profile_name.set('Meu rascunho')
        self.app.field_vars['font_size'].set('17')
        before = self.app._values()
        self.assertTrue(self.app.import_profile(self.source()))
        self.pump_until(lambda: self.app.import_dialog is not None)
        self.app.import_dialog.destroy()
        self.assertEqual(self.app._values(), before)
        self.assertEqual(self.app.profile_name.get(), 'Meu rascunho')
        self.assertTrue(self.app.dirty)
        self.assertFalse(self.app.config_model.stacks)

    def test_import_confirm_save_reload_and_name_collision(self):
        from docx.shared import Pt
        path = self.root / 'Modelo.docx'
        doc = Document()
        doc.add_paragraph('Texto modelo').runs[0].font.size = Pt(14)
        doc.add_heading('Título modelo', 1).runs[0].font.size = Pt(22)
        doc.save(path)
        original = path.read_bytes()
        self.app.profile_name.set('Modelo')
        self.assertTrue(self.app.save_profile())
        self.assertTrue(self.app.import_profile(path))
        self.pump_until(lambda: self.app.import_dialog is not None)
        dialog = self.app.import_dialog
        self.assertIn('14', dialog.summary.get('1.0', 'end'))
        dialog.confirm()
        self.assertFalse(dialog.winfo_exists())
        self.assertTrue(self.app.dirty)
        self.assertEqual(self.app.profile_name.get(), 'Modelo 2')
        self.assertNotIn('Modelo 2', self.app.store.load().stacks)
        self.assertEqual(self.app.field_vars['font_size'].get(), '14')
        self.assertTrue(self.app.save_profile())
        saved = self.app.store.load().stacks['Modelo 2']
        self.assertEqual(saved.font_size, 14)
        self.assertEqual(saved.category_rules['title']['values']['font_size'], 22)
        self.assertEqual(path.read_bytes(), original)

    def test_import_dirty_confirmation_can_be_cancelled(self):
        self.app.field_vars['font_size'].set('17')
        self.assertTrue(self.app.import_profile(self.source()))
        self.pump_until(lambda: self.app.import_dialog is not None)
        dialog = self.app.import_dialog
        with patch('app.ui.messagebox.askyesnocancel', return_value=None):
            dialog.confirm()
        self.assertTrue(dialog.winfo_exists())
        self.assertEqual(self.app.field_vars['font_size'].get(), '17')
        with patch('app.ui.messagebox.askyesnocancel', return_value=False):
            dialog.confirm()
        self.assertFalse(dialog.winfo_exists())
        self.assertTrue(self.app.dirty)

    def test_import_failure_picker_cancel_and_busy_leave_profile_intact(self):
        before = self.app._values()
        with patch('app.ui.filedialog.askopenfilename', return_value=''):
            self.assertFalse(self.app.import_profile())
        invalid = self.root / 'invalid.docx'
        invalid.write_text('not a Word')
        self.assertTrue(self.app.import_profile(invalid))
        self.pump_until(lambda: not self.app.busy)
        self.assertIn('importar', self.app.status_text.get())
        self.assertEqual(self.app._values(), before)
        self.assertIsNone(self.app.import_dialog)
        self.app._set_busy(True)
        self.assertFalse(self.app.import_profile(self.source()))
        self.assertEqual(self.app.import_button.cget('state'), 'disabled')
        self.app._set_busy(False)

    def test_import_page_and_category_alternatives_update_preview(self):
        from docx.shared import Cm, Pt
        path = self.root / 'Variações.docx'
        doc = Document()
        doc.add_paragraph('Primeiro padrão').runs[0].font.size = Pt(12)
        doc.add_paragraph('Segundo padrão').runs[0].font.size = Pt(16)
        doc.sections[0].left_margin = Cm(2)
        doc.add_section().left_margin = Cm(4)
        doc.save(path)
        self.assertTrue(self.app.import_profile(path))
        self.pump_until(lambda: self.app.import_dialog is not None)
        dialog = self.app.import_dialog
        dialog.page_selector.current(1)
        dialog.page_selector.event_generate('<<ComboboxSelected>>')
        dialog.option_selector.current(1)
        dialog.option_selector.event_generate('<<ComboboxSelected>>')
        self.app.update()
        expected = dialog.report.build_settings(page_index=1, selections={'body': 1})
        dialog.confirm()
        self.assertAlmostEqual(self.app._draft_base.margin_left_cm, expected.margin_left_cm)
        self.assertEqual(self.app._draft_base.font_size, expected.font_size)

    def test_import_cancel_during_read_discards_result_and_keeps_ui_responsive(self):
        import threading
        from app.profile_import import inspect_profile
        entered, release = threading.Event(), threading.Event()
        path = self.source()
        before = self.app._values()

        def delayed_read(source):
            entered.set()
            if not release.wait(5):
                raise RuntimeError('Test did not release import worker')
            return inspect_profile(source)

        with patch('app.profile_import.inspect_profile', side_effect=delayed_read):
            self.assertTrue(self.app.import_profile(path))
            self.pump_until(entered.is_set)
            self.assertTrue(self.app.busy)
            self.assertFalse(self.app.import_profile(path))
            self.app.cancel_batch()
            release.set()
            self.pump_until(lambda: not self.app.busy)
        self.assertIsNone(self.app.import_dialog)
        self.assertEqual(self.app._values(), before)
        self.assertIn('cancelada', self.app.status_text.get())

    def test_import_invalid_settings_disable_confirmation(self):
        from docx.shared import Cm
        path = self.root / 'invalid-geometry.docx'
        doc = Document()
        doc.add_paragraph('Um recuo que ultrapassa a página').paragraph_format.left_indent = Cm(40)
        doc.save(path)
        before = self.app._values()
        self.assertTrue(self.app.import_profile(path))
        self.pump_until(lambda: self.app.import_dialog is not None)
        dialog = self.app.import_dialog
        self.assertEqual(str(dialog.confirm_button.cget('state')), 'disabled')
        self.assertTrue(dialog.error_var.get())
        dialog.confirm()
        self.assertTrue(dialog.winfo_exists())
        self.assertEqual(self.app._values(), before)

    def test_import_close_during_read_finishes_without_opening_review(self):
        import threading
        from app.profile_import import inspect_profile
        release = threading.Event()
        path = self.source()

        def delayed_read(source):
            if not release.wait(5):
                raise RuntimeError('Test did not release import worker')
            return inspect_profile(source)

        with patch('app.profile_import.inspect_profile', side_effect=delayed_read):
            self.assertTrue(self.app.import_profile(path))
            with patch('app.ui.messagebox.askyesno', return_value=True):
                self.app.request_close()
            release.set()
            self.pump_until(lambda: self.app._destroyed)
        self.assertIsNone(self.app.import_dialog)
        self.assertFalse(self.app._after_ids)

    def test_save_reload_custom_font_decimal_and_active_profile(self):
        self.app.tabs.set('Perfis e formatação')
        self.app.profile_name.set('Meu perfil')
        for field, value in [('font_name', 'Fonte Personalizada'), ('font_size', '36,5'), ('line_spacing', '1,75')]:
            widget = self.app.field_widgets[field]
            if field == 'font_name':
                widget = widget._entry  # The editable entry inside CTkComboBox.
            widget.delete(0, 'end')
            widget.insert(0, value)
        self.assertTrue(self.app.save_profile())
        self.assertFalse(self.app.dirty)
        self.app.destroy()
        from app.ui import FormatWordApp
        self.app = FormatWordApp(config_dir=self.root / 'config')
        self.app.update()
        self.assertEqual(self.app.selected_profile.get(), 'Meu perfil')
        self.assertEqual(self.app.field_vars['font_name'].get(), 'Fonte Personalizada')
        self.assertEqual(self.app.field_vars['font_size'].get(), '36,5')
        self.assertEqual(self.app.field_vars['line_spacing'].get(), '1,75')

    def test_batch_invalid_then_valid_and_empty_destination(self):
        invalid = self.root / 'invalid.docx'
        invalid.write_text('not a word document')
        self.app.add_files([invalid, self.source()])
        self.app.output_dir.set('')
        self.assertFalse(self.app.start_batch())
        self.assertIn('destino', self.app.status_text.get().lower())
        self.app.output_dir.set(str(self.root / 'out'))
        self.assertTrue(self.app.start_batch())
        self.assertFalse(self.app.start_batch())
        self.pump_until(lambda: not self.app.busy)
        self.assertEqual(len(self.app.results), 2)
        self.assertIsNotNone(self.app.results[0].error)
        result = self.app.results[1].result
        self.assertTrue(result.output_path.is_file())
        self.assertEqual(Document(result.output_path).paragraphs[0].text, 'Conteúdo original')

    def test_empty_selection_and_invalid_form(self):
        self.app.output_dir.set(str(self.root / 'out'))
        self.assertFalse(self.app.start_batch())
        self.app.add_files([self.source()])
        self.app.field_vars['font_size'].set('abc')
        self.assertFalse(self.app.start_batch())
        self.assertIn('Tamanho da fonte', self.app.error_text.get())
        self.assertFalse(self.app.save_profile())

    def test_dirty_switch_cancel_and_transactional_save(self):
        self.app.profile_name.set('Salvo')
        self.assertTrue(self.app.save_profile())
        self.app.field_vars['font_size'].set('18')
        with patch('app.ui.messagebox.askyesnocancel', return_value=None):
            self.assertFalse(self.app.select_profile('Configuração atual'))
        self.assertEqual(self.app.selected_profile.get(), 'Salvo')
        self.assertEqual(self.app.field_vars['font_size'].get(), '18')
        previous = self.app.config_model
        with patch.object(self.app.store, 'save', side_effect=OSError('disco indisponível')):
            self.assertFalse(self.app.save_profile())
        self.assertIs(self.app.config_model, previous)
        self.assertTrue(self.app.dirty)

    def test_resize_and_conditional_fields(self):
        for width in (768, 1024, 1440):
            self.app.geometry(f'{width}x760')
            self.app.update()
            self.assertEqual(self.app._form_columns, 1 if width < 900 else 2)
            self.assertTrue(self.app.apply_button.winfo_exists())
        self.app.geometry('768x600')
        self.app.tabs.set('Arquivos e resultados')
        self.app.update()
        self.assertGreaterEqual(self.app.file_tree.winfo_height(), 60)
        for widget in (self.app.apply_button, self.app.cancel_button, self.app.status_label):
            bottom = widget.winfo_rooty() - self.app.winfo_rooty() + widget.winfo_height()
            self.assertLessEqual(bottom, self.app.winfo_height())
        self.assertEqual(self.app.field_widgets['page_width_cm'].cget('state'), 'normal')
        self.app.field_vars['paper_size'].set('Personalizado')
        self.app.update()
        self.assertEqual(self.app.field_widgets['page_width_cm'].cget('state'), 'normal')
        self.assertEqual(self.app.field_widgets['header_distance_cm'].cget('state'), 'normal')

    def test_keyboard_can_activate_primary_action(self):
        self.app.add_files([self.source()])
        self.app.output_dir.set(str(self.root / 'out'))
        canvas = self.app.apply_button._canvas
        self.assertEqual(str(canvas.cget('takefocus')), '1')
        self.assertGreaterEqual(self.app.apply_button.cget('height'), 36)
        self.app.focus_force()
        canvas.focus_force()
        self.app.update()
        canvas.event_generate('<Return>')
        self.app.update()
        self.pump_until(lambda: bool(self.app.results) and not self.app.busy)
        self.assertTrue(self.app.results[0].result.output_path.is_file())

    def test_apply_shortcut_does_not_also_invoke_focused_button(self):
        button = self.app.apply_button
        self.app.focus_force()
        button._canvas.focus_force()
        self.app.update()
        with patch.object(self.app, 'start_batch') as start:
            from unittest.mock import Mock
            other_action = Mock()
            button.configure(command=other_action)
            button._canvas.event_generate('<Control-Return>')
            self.app.update()
            start.assert_called_once_with()
            other_action.assert_not_called()

    def test_cancel_and_close_while_worker_active(self):
        # A gate controls timing, while the real batch processes real DOCX files.
        from app.batch import format_documents_batch
        import threading
        gate = threading.Event()
        def gated(*args, **kwargs):
            gate.wait(5)
            return format_documents_batch(*args, **kwargs)
        self.app.add_files([self.source('a.docx'), self.source('b.docx')])
        self.app.output_dir.set(str(self.root / 'out'))
        with patch('app.ui.format_documents_batch', side_effect=gated):
            self.assertTrue(self.app.start_batch())
            self.app.cancel_batch()
            self.assertTrue(self.app.cancel_event.is_set())
            self.assertTrue(self.app.busy)
            gate.set()
            self.pump_until(lambda: not self.app.busy)
        self.assertTrue(all(item.cancelled for item in self.app.results))
        gate.clear()
        with patch('app.ui.format_documents_batch', side_effect=gated):
            self.assertTrue(self.app.start_batch())
            with patch('app.ui.messagebox.askyesno', return_value=True):
                self.app.request_close()
            self.assertFalse(self.app._destroyed)
            gate.set()
            self.pump_until(lambda: self.app._destroyed)

    def test_duplicate_delete_and_restore_profiles(self):
        self.app.profile_name.set('Original')
        self.assertTrue(self.app.save_profile())
        self.app.duplicate_profile()
        self.assertEqual(set(self.app.config_model.stacks), {'Original', 'Original cópia'})
        self.app.field_vars['font_size'].set('20')
        with patch('app.ui.messagebox.askyesno', return_value=True):
            self.app.restore_profile()
        self.assertEqual(self.app.field_vars['font_size'].get(), '12')
        with patch('app.ui.messagebox.askyesno', return_value=True):
            self.app.delete_profile()
        self.assertEqual(set(self.app.config_model.stacks), {'Original'})
        self.assertEqual(self.app.selected_profile.get(), 'Configuração atual')

    def test_legacy_review_changes_draft_only_and_preserves_file(self):
        from app.config import AppConfig, FormatSettings
        template = self.root / 'old-template.docx'
        template.write_bytes(b'original asset')
        legacy = FormatSettings(template_path=str(template), migration_warnings=('Posição antiga: revisar',))
        self.app.config_model = AppConfig(legacy)
        self.app.store.save(self.app.config_model)
        self.app._load_draft(legacy, '')
        self.assertFalse(self.app.save_profile())
        with patch('app.ui.messagebox.askyesno', return_value=True):
            self.app.review_legacy()
        self.assertTrue(self.app.dirty)
        self.assertEqual(self.app._draft_base.template_path, '')
        self.assertEqual(self.app.config_model.settings.template_path, str(template))
        self.assertEqual(template.read_bytes(), b'original asset')
        self.assertTrue(self.app.save_profile())
        self.assertEqual(self.app.store.load().settings.template_path, '')

    def test_inactive_invalid_numbers_remain_editable_until_corrected(self):
        cases = (
            ('paper_size', 'Personalizado', 'A4', 'page_width_cm', '21'),
            ('header_mode', 'Imagem do perfil', 'Preservar original', 'header_image_width_cm', '15'),
            ('footer_mode', 'Imagem do perfil', 'Preservar original', 'footer_image_width_cm', '15'),
        )
        for mode, active, inactive, field, corrected in cases:
            with self.subTest(field=field):
                self.app.field_vars[mode].set(active)
                widget = self.app.field_widgets[field]
                widget.delete(0, 'end')
                self.app.field_vars[mode].set(inactive)
                self.app.update()
                self.assertEqual(self.app.field_vars[field].get(), '')
                self.assertEqual(widget.cget('state'), 'normal')
                self.assertFalse(self.app.save_profile())
                self.assertTrue(self.app.field_errors[field].cget('text'))
                self.assertEqual(self.app.field_errors[field].cget('text'), self.app.error_text.get())
                widget.insert(0, corrected)
                self.assertEqual(self.app.field_vars[field].get(), corrected)
                self.assertTrue(self.app.save_profile())

    def test_category_rules_confirm_cancel_restore_and_deep_isolation(self):
        self.app.new_profile()
        self.assertEqual(self.app.field_vars['formatting_mode'].get(), 'Por tipo de texto')
        self.assertTrue(self.app.save_profile())
        dialog = self.app.edit_category_rules()
        dialog.category_var.set('Título')
        dialog.switch_category()
        dialog.mode_var.set('Personalizar')
        dialog.variables['font_size'].set('22')
        dialog.confirm()
        self.assertTrue(self.app.dirty)
        self.assertEqual(self.app._draft_base.category_rules['title']['values']['font_size'], 22)
        self.assertNotIn('title', self.app._baseline_base.category_rules)
        self.assertTrue(self.app.save_profile())
        saved = self.app.config_model.stacks[self.app._selected]
        self.app._draft_base.category_rules['title']['values']['font_size'] = 30
        self.assertEqual(saved.category_rules['title']['values']['font_size'], 22)
        self.assertEqual(self.app._baseline_base.category_rules['title']['values']['font_size'], 22)
        with patch('app.ui.messagebox.askyesno', return_value=True):
            self.app.restore_profile()
        self.assertEqual(self.app._draft_base.category_rules['title']['values']['font_size'], 22)
        dialog = self.app.edit_category_rules()
        dialog.settings.category_rules['title']['values']['font_size'] = 40
        dialog.destroy()
        self.assertFalse(self.app.dirty)
        self.assertEqual(saved.category_rules['title']['values']['font_size'], 22)

    def test_review_is_transactional_retained_and_sent_to_batch(self):
        path = self.source()
        self.app.add_files([path])
        self.app.file_tree.selection_set('0')
        dialog = self.app.review_document()
        dialog.paragraph_tree.selection_set('0')
        dialog.category_var.set('Título')
        dialog.same_style.set(True)
        dialog.classify_selected()
        self.assertNotIn(path, self.app.document_overrides)
        self.assertNotIn('Normal', self.app._draft_base.style_categories)
        dialog.confirm()
        self.assertEqual(self.app.document_overrides[path].paragraph_categories, {0: 'title'})
        self.assertEqual(self.app._draft_base.style_categories['Normal'], 'title')
        self.assertNotIn('Normal', self.app._baseline_base.style_categories)
        self.app.add_files([self.source('other.docx')])
        self.assertIn(path, self.app.document_overrides)
        self.app.output_dir.set(str(self.root / 'out'))
        with patch('app.ui.format_documents_batch') as batch:
            self.assertTrue(self.app.start_batch())
            self.pump_until(lambda: not self.app.busy)
        snapshot = batch.call_args.kwargs['overrides_by_path'][path]
        self.app.document_overrides[path].paragraph_categories[0] = 'body'
        self.assertEqual(snapshot.paragraph_categories[0], 'title')
        self.app.clear_files()
        self.assertEqual(self.app.document_overrides, {})

    def test_review_image_validation_and_confirm(self):
        from PIL import Image
        from docx.shared import Cm
        image = self.root / 'picture.png'
        Image.new('RGB', (80, 40), 'red').save(image)
        path = self.root / 'images.docx'
        doc = Document()
        doc.add_paragraph('Destino da imagem')
        doc.add_picture(str(image), width=Cm(3))
        doc.save(path)
        self.app.add_files([path])
        self.app.file_tree.selection_set('0')
        dialog = self.app.review_document()
        dialog.image_tree.selection_set('0')
        dialog.select_image()
        dialog.width_var.set('nan')
        dialog.confirm()
        self.assertTrue(dialog.winfo_exists())
        self.assertNotIn(path, self.app.document_overrides)
        dialog.width_var.set('4,5')
        dialog.alignment_var.set('Centralizado')
        dialog.target_var.set(next(label for label, value in dialog.targets.items() if value == 0))
        dialog.placement_var.set('Antes')
        dialog.confirm()
        edit = self.app.document_overrides[path].image_edits[0]
        self.assertEqual((edit.width_cm, edit.alignment, edit.target_paragraph, edit.placement), (4.5, 'center', 0, 'before'))

    def test_category_editor_validates_custom_categories_independently(self):
        from app.config import FormatSettings, validate_settings
        settings = FormatSettings(formatting_mode='by_category', category_rules={
            'body': {'mode': 'custom', 'values': {'left_indent_cm': 10}},
            'title': {'mode': 'custom', 'values': {'right_indent_cm': 10}},
        })
        self.app._load_draft(validate_settings(settings), '')
        dialog = self.app.edit_category_rules()
        dialog.category_var.set('Título')
        dialog.switch_category()
        self.assertEqual(dialog.current_category, 'title', dialog.error_var.get())
        dialog.variables['font_size'].set('18')
        dialog.confirm()
        self.assertFalse(dialog.winfo_exists())
        rules = self.app._draft_base.category_rules
        self.assertEqual(rules['body']['values']['left_indent_cm'], 10)
        self.assertEqual(rules['title']['values']['right_indent_cm'], 10)
        self.assertEqual(rules['title']['values']['font_size'], 18)
        self.assertTrue(self.app.save_profile())

    def test_review_blank_width_preserves_original_over_profile_image_rule(self):
        from PIL import Image
        from docx.shared import Cm
        image = self.root / 'original.png'
        Image.new('RGB', (80, 40), 'red').save(image)
        path = self.root / 'preserve-image.docx'
        doc = Document()
        doc.add_picture(str(image), width=Cm(3))
        doc.save(path)
        self.app.field_vars['body_image_width_cm'].set('8')
        self.app.field_vars['body_image_alignment'].set('Centralizado')
        self.app.add_files([path])
        self.app.file_tree.selection_set('0')
        dialog = self.app.review_document()
        dialog.image_tree.selection_set('0')
        dialog.select_image()
        self.assertEqual(dialog.width_var.get(), '8')
        self.assertEqual(dialog.alignment_var.get(), 'Centralizado')
        dialog.width_var.set('')
        dialog.alignment_var.set('Preservar')
        dialog.confirm()
        edit = self.app.document_overrides[path].image_edits[0]
        self.assertIsNone(edit.width_cm)
        self.assertEqual(edit.alignment, 'preserve')
        self.app.output_dir.set(str(self.root / 'preserved'))
        self.assertTrue(self.app.start_batch())
        self.pump_until(lambda: not self.app.busy)
        self.assertIsNone(self.app.results[0].error)
        result = Document(self.app.results[0].result.output_path)
        self.assertAlmostEqual(result.inline_shapes[0].width.cm, 3, places=4)
