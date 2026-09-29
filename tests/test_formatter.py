from __future__ import annotations

import hashlib
import errno
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

from app.config import FormatSettings
from app.formatter import FormatterError, format_document


class FormatterTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'entrada.docx'
        self.outdir = self.root / 'output'
        self.doc = Document()
        self.doc.add_paragraph('Primeiro texto').runs[0].bold = True
        self.image = self.root / 'image.png'
        Image.new('RGB', (800, 80), 'navy').save(self.image)

    def run_format(self, settings=None):
        self.doc.save(self.source)
        result = format_document(self.source, self.outdir, settings or FormatSettings())
        return Document(result.output_path), result

    def test_preserves_structure_and_source_bytes(self):
        self.doc.paragraphs[0].style = 'Heading 1'
        table = self.doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Tabela'
        table.cell(1, 0).merge(table.cell(1, 1)).text = 'Mesclada'
        table.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = 'Aninhada'
        self.doc.add_paragraph('Depois da tabela', style='List Bullet')
        self.doc.add_picture(str(self.image))
        self.doc.add_section()
        self.doc.add_paragraph('Seção dois')
        self.doc.save(self.source)
        original = self.source.read_bytes()
        result = format_document(self.source, self.outdir, FormatSettings())
        out = Document(result.output_path)
        self.assertEqual(len(out.tables), 1)
        self.assertEqual(len(out.tables[0].cell(0, 1).tables), 1)
        self.assertEqual(len(out.inline_shapes), 1)
        self.assertEqual(len(out.sections), 2)
        self.assertEqual(out.paragraphs[0].style.name, 'Heading 1')
        self.assertTrue(out.paragraphs[0].runs[0].bold)
        self.assertEqual([e.tag for e in out._element.body], [e.tag for e in self.doc._element.body])
        self.assertEqual(self.source.read_bytes(), original)

    def test_preserves_spaces_empty_paragraphs_tabs_and_breaks(self):
        self.doc.add_paragraph('  Espaços  e\ttab\nquebra ')
        self.doc.add_paragraph('')
        self.doc.add_page_break()
        out, _ = self.run_format()
        self.assertEqual([p.text for p in out.paragraphs], [p.text for p in self.doc.paragraphs])
        self.assertEqual(len(out._element.xpath('.//w:br')), 2)

    def test_formats_hyperlink_without_losing_relationship(self):
        p = self.doc.paragraphs[0]
        link = OxmlElement('w:hyperlink')
        rel_id = p.part.relate_to('https://example.com', RT.HYPERLINK, is_external=True)
        link.set(qn('r:id'), rel_id)
        run = OxmlElement('w:r')
        text = OxmlElement('w:t')
        text.text = 'Visitar'
        run.append(text)
        link.append(run)
        p._p.append(link)
        out, _ = self.run_format(FormatSettings(font_size=36))
        r = out.paragraphs[0]._p.xpath('./w:hyperlink/w:r')[0]
        self.assertEqual(r.rPr.sz.val.pt, 36)
        self.assertEqual(out.part.rels[rel_id].target_ref, 'https://example.com')

    def test_font_decimal_and_theme_overrides(self):
        rpr = self.doc.paragraphs[0].runs[0]._r.get_or_add_rPr()
        fonts = rpr.get_or_add_rFonts()
        fonts.set(qn('w:asciiTheme'), 'majorHAnsi')
        fonts.set(qn('w:eastAsiaTheme'), 'majorEastAsia')
        out, _ = self.run_format(FormatSettings(font_name='Verdana', font_size=12.5))
        run = out.paragraphs[0].runs[0]
        self.assertEqual(run.font.size.pt, 12.5)
        for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            self.assertEqual(run._r.rPr.rFonts.get(qn('w:' + attr)), 'Verdana')
        self.assertIsNone(run._r.rPr.rFonts.get(qn('w:asciiTheme')))
        self.assertIsNone(run._r.rPr.rFonts.get(qn('w:eastAsiaTheme')))
        self.assertEqual(run._r.rPr.find(qn('w:szCs')).get(qn('w:val')), '25')

    def test_all_paragraph_options_in_nested_tables(self):
        table = self.doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'Célula'
        table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = 'Outra'
        settings = FormatSettings(alignment='right', paragraph_spacing_before=3.5,
            paragraph_spacing_after=8.5, line_spacing=2, left_indent_cm=.4,
            right_indent_cm=.6, first_line_indent_cm=-.3,
            keep_with_next=True, keep_together=False, widow_control=True)
        out, _ = self.run_format(settings)
        from docx.text.paragraph import Paragraph
        for element in out._element.body.xpath('.//w:p'):
            p = Paragraph(element, out._body)
            f = p.paragraph_format
            self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.RIGHT)
            self.assertAlmostEqual(f.space_before.pt, 3.5)
            self.assertAlmostEqual(f.space_after.pt, 8.5)
            self.assertAlmostEqual(f.left_indent.cm, .4, places=2)
            self.assertAlmostEqual(f.right_indent.cm, .6, places=2)
            self.assertAlmostEqual(f.first_line_indent.cm, -.3, places=2)
            self.assertEqual(f.line_spacing, 2)
            self.assertTrue(f.keep_with_next)
            self.assertFalse(f.keep_together)
            self.assertTrue(f.widow_control)

    def test_automatic_spacing_and_character_indents_do_not_override_settings(self):
        ppr = self.doc.paragraphs[0]._p.get_or_add_pPr()
        spacing = ppr.get_or_add_spacing()
        for key in ('beforeAutospacing', 'afterAutospacing', 'beforeLines', 'afterLines'):
            spacing.set(qn('w:' + key), '100')
        ind = ppr.get_or_add_ind()
        ind.set(qn('w:firstLineChars'), '500')
        ind.set(qn('w:start'), '900')
        out, _ = self.run_format()
        ppr = out.paragraphs[0]._p.pPr
        self.assertNotIn(qn('w:beforeLines'), ppr.spacing.attrib)
        self.assertNotIn(qn('w:firstLineChars'), ppr.ind.attrib)
        self.assertNotIn(qn('w:start'), ppr.ind.attrib)
        self.assertEqual(ppr.spacing.get(qn('w:beforeAutospacing')), '0')
        self.assertEqual(ppr.spacing.get(qn('w:afterAutospacing')), '0')

    def test_exact_and_minimum_line_spacing(self):
        for mode, expected in [('exact', WD_LINE_SPACING.EXACTLY), ('at_least', WD_LINE_SPACING.AT_LEAST)]:
            with self.subTest(mode=mode):
                out, _ = self.run_format(FormatSettings(line_spacing_mode=mode, line_spacing=17.5))
                f = out.paragraphs[0].paragraph_format
                self.assertEqual(f.line_spacing_rule, expected)
                self.assertEqual(f.line_spacing.pt, 17.5)

    def test_alignment_choices(self):
        for name, value in [('left', 0), ('center', 1), ('right', 2), ('justify', 3)]:
            with self.subTest(name=name):
                out, _ = self.run_format(FormatSettings(alignment=name))
                self.assertEqual(out.paragraphs[0].alignment, value)

    def test_explicit_font_emphasis_and_color(self):
        run = self.doc.paragraphs[0].runs[0]
        run.italic = True
        run.underline = True
        out, _ = self.run_format(FormatSettings(bold=False, italic=False, underline=False, font_color='#12ab34'))
        r = out.paragraphs[0].runs[0]
        self.assertFalse(r.bold)
        self.assertFalse(r.italic)
        self.assertFalse(r.underline)
        self.assertEqual(str(r.font.color.rgb), '12AB34')

    def test_enable_emphasis_overrides_disabled_style(self):
        out, _ = self.run_format(FormatSettings(bold=True, italic=True, underline=True))
        run = out.paragraphs[0].runs[0]
        self.assertTrue(run.bold)
        self.assertTrue(run.italic)
        self.assertTrue(run.underline)

    def test_all_sections_receive_page_configuration(self):
        self.doc.add_section()
        out, _ = self.run_format(FormatSettings(paper_size='Letter', orientation='landscape',
            margin_top_cm=2.1, margin_bottom_cm=2.2, margin_left_cm=2.3, margin_right_cm=2.4,
            header_distance_cm=.6, footer_distance_cm=.5))
        for section in out.sections:
            self.assertAlmostEqual(section.page_width.cm, 27.94, places=2)
            self.assertAlmostEqual(section.page_height.cm, 21.59, places=2)
            for key, expected in [('top_margin',2.1),('bottom_margin',2.2),('left_margin',2.3),('right_margin',2.4),('header_distance',.6),('footer_distance',.5)]:
                self.assertAlmostEqual(getattr(section,key).cm, expected, places=2)

    def test_custom_page_size(self):
        out, _ = self.run_format(FormatSettings(paper_size='custom', page_width_cm=18, page_height_cm=25))
        self.assertAlmostEqual(out.sections[0].page_width.cm, 18, places=2)
        self.assertAlmostEqual(out.sections[0].page_height.cm, 25, places=2)

    def test_preserves_all_header_footer_variants_by_default(self):
        section = self.doc.sections[0]
        section.different_first_page_header_footer = True
        self.doc.settings.odd_and_even_pages_header_footer = True
        for attr in ('header','first_page_header','even_page_header','footer','first_page_footer','even_page_footer'):
            getattr(section,attr).paragraphs[0].text = attr
        out, _ = self.run_format()
        for attr in ('header','first_page_header','even_page_header','footer','first_page_footer','even_page_footer'):
            self.assertEqual(getattr(out.sections[0], attr).paragraphs[0].text, attr)

    def test_remove_headers_and_footers_in_all_sections(self):
        self.doc.sections[0].header.paragraphs[0].text = 'Cabeçalho'
        self.doc.sections[0].first_page_footer.paragraphs[0].text = 'Rodapé'
        self.doc.add_section()
        self.doc.sections[1].header.is_linked_to_previous = False
        self.doc.sections[1].header.paragraphs[0].text = 'Segundo'
        out, result = self.run_format(FormatSettings(header_mode='remove',footer_mode='remove'))
        for section in out.sections:
            self.assertEqual(section._sectPr.xpath('./w:headerReference|./w:footerReference'), [])
        with ZipFile(result.output_path) as archive:
            self.assertFalse([name for name in archive.namelist() if name.startswith(('word/header', 'word/footer'))])

    def test_image_size_proportions_and_alignment(self):
        settings = FormatSettings(header_mode='image', header_image_path=str(self.image),
            header_image_width_cm=12, header_alignment='right', footer_mode='image',
            footer_image_path=str(self.image), footer_image_width_cm=10,
            footer_alignment='left', margin_bottom_cm=2.5)
        out, _ = self.run_format(settings)
        for slot, width, alignment in [('header',12,WD_ALIGN_PARAGRAPH.RIGHT),('footer',10,WD_ALIGN_PARAGRAPH.LEFT)]:
            p = getattr(out.sections[0],slot).paragraphs[0]
            extent = p._p.xpath('.//wp:extent')[0]
            self.assertAlmostEqual(int(extent.get('cx')) / 360000, width, places=2)
            self.assertAlmostEqual(int(extent.get('cy')) / int(extent.get('cx')), .1, places=3)
            self.assertEqual(p.alignment, alignment)
            self.assertEqual(p.paragraph_format.first_line_indent, 0)

    def test_missing_image_is_error_not_silent_success(self):
        self.doc.save(self.source)
        with self.assertRaises(FormatterError):
            format_document(self.source, self.outdir, FormatSettings(header_mode='image',header_image_path=str(self.root/'missing.png')))
        self.assertFalse(list(self.outdir.glob('*.docx')))

    def test_image_that_does_not_fit_margins_is_rejected(self):
        self.doc.save(self.source)
        with self.assertRaisesRegex(FormatterError, 'cabeçalho|Cabeçalho'):
            format_document(self.source, self.outdir, FormatSettings(header_mode='image',header_image_path=str(self.image),header_image_width_cm=15,margin_top_cm=1))

    def test_legacy_template_is_error_not_silent_success(self):
        self.doc.save(self.source)
        with self.assertRaisesRegex(FormatterError, '[Tt]emplate'):
            format_document(self.source,self.outdir,FormatSettings(template_path=str(self.root/'missing.docx')))

    def test_invalid_numeric_is_rejected_without_output(self):
        self.doc.save(self.source)
        for value in [float('nan'), float('inf'), 0, -1, 'abc', True]:
            with self.subTest(value=value), self.assertRaises(FormatterError):
                format_document(self.source,self.outdir,FormatSettings(font_size=value))
        self.assertFalse(list(self.outdir.glob('*.docx')))

    def test_empty_and_image_only_documents_are_supported(self):
        self.doc = Document()
        out, _ = self.run_format()
        self.assertEqual(len(out.paragraphs), 0)
        self.doc.add_picture(str(self.image))
        out, _ = self.run_format()
        self.assertEqual(len(out.inline_shapes), 1)

    def test_rejects_fake_docx_and_pdf(self):
        for suffix in ['.docx','.pdf','.doc']:
            path = self.root / ('fake' + suffix)
            path.write_text('invalid')
            with self.subTest(suffix=suffix), self.assertRaises(FormatterError):
                format_document(path,self.outdir,FormatSettings())

    def test_rejects_tracked_changes_and_textboxes(self):
        for tag in ['w:ins', 'w:txbxContent', 'w:altChunk', 'm:oMath']:
            self.doc = Document()
            self.doc.add_paragraph('Conteúdo')
            self.doc._element.body.insert(0,OxmlElement(tag))
            self.doc.save(self.source)
            with self.subTest(tag=tag), self.assertRaises(FormatterError):
                format_document(self.source,self.outdir,FormatSettings(complex_content_mode='reject'))

    def test_fields_are_preserved_with_explicit_warning(self):
        run = self.doc.paragraphs[0].add_run()
        fld = OxmlElement('w:fldChar')
        fld.set(qn('w:fldCharType'),'begin')
        run._r.append(fld)
        instruction = OxmlElement('w:instrText')
        instruction.text = ' PAGE '
        run._r.append(instruction)
        out, result = self.run_format()
        self.assertEqual(out._element.xpath('.//w:instrText')[0].text,' PAGE ')
        self.assertTrue(result.warnings)

    def test_existing_outputs_and_input_not_overwritten(self):
        self.doc.save(self.source)
        before = hashlib.sha256(self.source.read_bytes()).digest()
        first = format_document(self.source,self.root,FormatSettings())
        snapshot = first.output_path.read_bytes()
        second = format_document(self.source,self.root,FormatSettings(font_size=14))
        self.assertNotEqual(first.output_path,second.output_path)
        self.assertEqual(first.output_path.read_bytes(),snapshot)
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).digest(),before)

    def test_simultaneous_exports_reserve_different_names(self):
        self.doc.save(self.source)
        with ThreadPoolExecutor(max_workers=3) as pool:
            results = list(pool.map(lambda _:format_document(self.source,self.outdir,FormatSettings()),range(3)))
        self.assertEqual(len({r.output_path for r in results}),3)
        for result in results:
            self.assertEqual(Document(result.output_path).paragraphs[0].text,'Primeiro texto')

    def test_save_failure_leaves_no_final_or_temporary_file(self):
        self.doc.save(self.source)
        with patch('docx.document.Document.save',side_effect=OSError('disco cheio')):
            with self.assertRaises(FormatterError):
                format_document(self.source,self.outdir,FormatSettings())
        self.assertEqual(list(self.outdir.iterdir()),[])

    def test_rejects_output_directory_that_is_file(self):
        self.doc.save(self.source)
        self.outdir.write_text('file')
        with self.assertRaises(FormatterError):
            format_document(self.source,self.outdir,FormatSettings())

    def test_no_new_header_parts_when_preserving(self):
        out, result = self.run_format()
        with ZipFile(result.output_path) as archive:
            self.assertFalse([name for name in archive.namelist() if name.startswith('word/header')])

    def test_output_without_requested_formatting_is_never_published(self):
        self.doc.save(self.source)
        original = self.source.read_bytes()
        with patch('docx.document.Document.save', side_effect=lambda stream: stream.write(original)):
            with self.assertRaisesRegex(FormatterError, 'verifica|Verifica'):
                format_document(self.source, self.outdir, FormatSettings(font_size=36))
        self.assertEqual(list(self.outdir.iterdir()), [])

    def test_export_to_filesystem_without_hardlinks(self):
        self.doc.save(self.source)
        with patch('app.formatter.os.link', side_effect=OSError(errno.ENOTSUP, 'unsupported')):
            first = format_document(self.source, self.outdir, FormatSettings())
            second = format_document(self.source, self.outdir, FormatSettings())
        self.assertNotEqual(first.output_path, second.output_path)
        self.assertEqual(Document(first.output_path).paragraphs[0].text, 'Primeiro texto')

    def test_new_properties_follow_word_schema_order(self):
        self.doc.add_section()
        out, _ = self.run_format(FormatSettings(bold=False, italic=False))
        for p in out.paragraphs:
            names = [e.tag.rsplit('}', 1)[-1] for e in p._p.pPr]
            self.assertLess(names.index('snapToGrid'), names.index('spacing'))
            self.assertLess(names.index('contextualSpacing'), names.index('jc'))
            if 'sectPr' in names:
                self.assertEqual(names[-1], 'sectPr')
        names = [e.tag.rsplit('}', 1)[-1] for e in out.paragraphs[0].runs[0]._r.rPr]
        self.assertLess(names.index('bCs'), names.index('sz'))
        self.assertLess(names.index('iCs'), names.index('sz'))

    def test_image_dpi_does_not_distort_proportions(self):
        Image.new('RGB', (800, 80), 'navy').save(self.image, dpi=(300, 72))
        out, _ = self.run_format(FormatSettings(header_mode='image', header_image_path=str(self.image), header_image_width_cm=12))
        extent = out.sections[0].header.paragraphs[0]._p.xpath('.//wp:extent')[0]
        self.assertAlmostEqual(int(extent.get('cy')) / int(extent.get('cx')), .1, places=3)

    def test_empty_paragraph_mark_receives_font_and_size(self):
        self.doc.add_paragraph('')
        out, _ = self.run_format(FormatSettings(font_name='Verdana', font_size=36))
        for p in out.paragraphs:
            mark = p._p.xpath('./w:pPr/w:rPr')
            self.assertEqual(len(mark), 1)
            self.assertEqual(mark[0].sz.val.pt, 36)
            self.assertEqual(mark[0].rFonts.get(qn('w:ascii')), 'Verdana')

    def test_header_image_does_not_inherit_giant_paragraph_mark(self):
        from docx.shared import Pt
        self.doc.styles['Normal'].font.size = Pt(100)
        out, _ = self.run_format(FormatSettings(header_mode='image', header_image_path=str(self.image), header_image_width_cm=10))
        p = out.sections[0].header.paragraphs[0]
        marks = p._p.xpath('./w:pPr/w:rPr/w:sz')
        self.assertEqual(len(marks), 1)
        self.assertEqual(marks[0].val.pt, 1)
        self.assertEqual(p._p.pPr.spacing.get(qn('w:beforeAutospacing')), '0')


if __name__ == '__main__':
    unittest.main()
