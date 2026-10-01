import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.config import SettingsError, TEXT_FIELDS
from app.document_model import fingerprint
from app.formatter import FormatterError, format_document


class ProfileImportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / 'reference.docx'
        self.doc = Document()

    def inspect(self):
        from app.profile_import import inspect_profile
        self.doc.save(self.source)
        return inspect_profile(self.source)

    def test_direct_properties_and_immutable_source(self):
        p = self.doc.add_paragraph('Texto de referência')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(8)
        r = p.runs[0]
        r.font.name = 'Georgia'; r.font.size = Pt(13)
        r.bold = True; r.italic = False; r.underline = True
        r.font.color.rgb = RGBColor.from_string('123456')
        self.doc.sections[0].left_margin = Cm(3)
        self.doc.save(self.source)
        before = fingerprint(self.source)
        from app.profile_import import inspect_profile
        report = inspect_profile(self.source)
        settings = report.build_settings()
        self.assertEqual(fingerprint(self.source), before)
        self.assertEqual(settings.font_name, 'Georgia')
        self.assertEqual(settings.font_size, 13)
        self.assertEqual(settings.font_color, '123456')
        self.assertTrue(settings.bold)
        self.assertTrue(settings.underline)
        self.assertEqual(settings.alignment, 'justify')
        self.assertAlmostEqual(settings.first_line_indent_cm, 1, places=2)
        self.assertAlmostEqual(settings.margin_left_cm, 3, places=2)
        self.assertEqual(settings.paragraph_spacing_after, 8)
        self.assertEqual(set(report.category_options['body'][0].values), set(TEXT_FIELDS))
        self.assertEqual(settings.formatting_mode, 'by_category')
        self.assertEqual(settings.category_rules['body']['mode'], 'inherit')
        self.assertEqual(settings.category_rules['quote']['mode'], 'preserve')

    def test_inherited_paragraph_and_character_styles(self):
        base = self.doc.styles.add_style('Base legal', WD_STYLE_TYPE.PARAGRAPH)
        base.font.name = 'Verdana'; base.font.size = Pt(14)
        base.paragraph_format.left_indent = Cm(1)
        base.paragraph_format.keep_with_next = True
        child = self.doc.styles.add_style('Filho', WD_STYLE_TYPE.PARAGRAPH)
        child.base_style = base
        char = self.doc.styles.add_style('Destaque', WD_STYLE_TYPE.CHARACTER)
        char.font.italic = True
        p = self.doc.add_paragraph('Herdado', child)
        p.runs[0].style = char
        settings = self.inspect().build_settings()
        self.assertEqual(settings.font_name, 'Verdana')
        self.assertEqual(settings.font_size, 14)
        self.assertTrue(settings.italic)
        self.assertTrue(settings.keep_with_next)
        self.assertAlmostEqual(settings.left_indent_cm, 1, places=2)

    def test_defaults_theme_fonts_and_colors(self):
        p = self.doc.add_paragraph('Tema')
        rpr = p.runs[0]._r.get_or_add_rPr()
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '000000'); color.set(qn('w:themeColor'), 'accent1')
        rpr.append(color)
        settings = self.inspect().build_settings()
        self.assertEqual(settings.font_name, 'Cambria')
        self.assertEqual(settings.font_size, 11)
        self.assertEqual(settings.font_color, '4F81BD')
        self.assertAlmostEqual(settings.line_spacing, 1.15)

    def test_coherent_alternatives_by_category_with_inline_emphasis(self):
        for text in ('Principal texto comum', 'Segundo texto comum'):
            p = self.doc.add_paragraph(text)
            p.runs[0].font.name = 'Arial'; p.runs[0].font.size = Pt(12)
            p.add_run(' destaque').bold = True
        other = self.doc.add_paragraph('Outra opção')
        other.alignment = WD_ALIGN_PARAGRAPH.CENTER
        other.runs[0].font.name = 'Georgia'; other.runs[0].font.size = Pt(16)
        self.doc.add_heading('Título', 1)
        report = self.inspect()
        settings = report.build_settings()
        self.assertEqual(settings.font_name, 'Arial')
        self.assertIsNone(settings.bold)
        self.assertEqual(report.category_options['body'][0].count, 2)
        self.assertEqual(settings.category_rules['title']['mode'], 'custom')
        index = next(i for i, o in enumerate(report.category_options['body']) if o.values['font_name'] == 'Georgia')
        alternate = report.build_settings(selections={'body': index})
        self.assertEqual(alternate.font_size, 16)
        self.assertEqual(alternate.alignment, 'center')
        self.assertTrue(any('ênfase mista' in n for n in report.notices))

    def test_sections_are_individually_selectable(self):
        self.doc.add_paragraph('Corpo')
        first = self.doc.sections[0]
        first.page_width = Cm(21); first.page_height = Cm(29.7)
        first.left_margin = Cm(3)
        second = self.doc.add_section(WD_SECTION.NEW_PAGE)
        second.page_width = Cm(29.7); second.page_height = Cm(21)
        second.left_margin = Cm(2)
        report = self.inspect()
        self.assertEqual(len(report.page_options), 2)
        self.assertEqual(report.build_settings(page_index=1).orientation, 'landscape')
        self.assertAlmostEqual(report.build_settings(page_index=1).page_width_cm, 21, places=2)
        self.assertAlmostEqual(report.build_settings(page_index=1).page_height_cm, 29.7, places=2)
        self.assertAlmostEqual(report.build_settings(page_index=0).margin_left_cm, 3, places=2)

    def test_invalid_selected_values_not_clamped(self):
        p = self.doc.add_paragraph('Muito pequeno')
        p.paragraph_format.line_spacing = .25
        report = self.inspect()
        self.assertEqual(report.category_options['body'][0].values['line_spacing'], .25)
        with self.assertRaisesRegex(SettingsError, 'alternativa importada'):
            report.build_settings()
        with self.assertRaises(SettingsError):
            report.build_settings(page_index=-1)
        with self.assertRaises(SettingsError):
            report.build_settings(selections={'body': 999})

    def test_invalid_extension_archive_and_input_limit(self):
        from app.profile_import import inspect_profile
        invalid = self.root / 'invalid.docx'; invalid.write_text('not a zip')
        with self.assertRaises(FormatterError):
            inspect_profile(invalid)
        invalid = self.root / 'invalid.doc'; invalid.write_text('not a word')
        with self.assertRaisesRegex(FormatterError, '.docx'):
            inspect_profile(invalid)
        self.doc.save(self.source)
        with self.assertRaises(SettingsError):
            inspect_profile(self.source, max_input_mb=0)

    def test_applying_imported_profile_preserves_destination_header(self):
        p = self.doc.add_paragraph('Referência')
        p.runs[0].font.name = 'Georgia'; p.runs[0].font.size = Pt(15)
        self.doc.sections[0].header.paragraphs[0].text = 'Timbrado referência'
        self.doc.add_heading('Referência título', 1)
        report = self.inspect()
        destination = Document()
        destination.add_paragraph('Destino')
        destination.add_heading('Destino título', 1)
        destination.sections[0].header.paragraphs[0].text = 'Timbrado destino'
        target = self.root / 'target.docx'; destination.save(target)
        result = format_document(target, self.root / 'out', report.build_settings())
        out = Document(result.output_path)
        self.assertEqual(out.paragraphs[0].runs[0].font.name, 'Georgia')
        self.assertEqual(out.paragraphs[0].runs[0].font.size.pt, 15)
        self.assertEqual(out.sections[0].header.paragraphs[0].text, 'Timbrado destino')
        self.assertTrue(any('NÃO são copiados' in n for n in report.notices))

    def test_unresolved_font_is_explicit_and_no_body_preserves(self):
        defaults = self.doc.styles.element.find(qn('w:docDefaults'))
        defaults.getparent().remove(defaults)
        self.doc.add_paragraph('Sem padrão')
        report = self.inspect()
        self.assertTrue(any('Nome da fonte não resolvido' in n for n in report.notices))
        self.assertEqual(report.build_settings().font_name, 'Arial')
        self.doc = Document()
        self.doc.add_heading('Somente título', 1)
        report = self.inspect()
        self.assertEqual(report.build_settings().category_rules['body']['mode'], 'preserve')

    def test_property_level_spacing_inheritance(self):
        base = self.doc.styles.add_style('Entrelinha exata', WD_STYLE_TYPE.PARAGRAPH)
        base.paragraph_format.line_spacing = Pt(18)
        p = self.doc.add_paragraph('Herde a regra', base)
        spacing = OxmlElement('w:spacing'); spacing.set(qn('w:line'), '400')
        p._p.get_or_add_pPr().append(spacing)
        settings = self.inspect().build_settings()
        self.assertEqual(settings.line_spacing_mode, 'exact')
        self.assertEqual(settings.line_spacing, 20)

    def test_unknown_theme_font_uses_disclosed_fallback(self):
        self.doc.styles['Normal'].font.name = 'Georgia'
        p = self.doc.add_paragraph('Tema desconhecido')
        fonts = OxmlElement('w:rFonts')
        fonts.set(qn('w:asciiTheme'), 'notAThemeSlot')
        p.runs[0]._r.get_or_add_rPr().append(fonts)
        report = self.inspect()
        self.assertEqual(report.build_settings().font_name, 'Arial')
        self.assertTrue(any('Fonte de tema não resolvida' in n for n in report.notices))

    def test_styles_toggle_bold_but_direct_false_clears_it(self):
        base = self.doc.styles.add_style('Base bold', WD_STYLE_TYPE.PARAGRAPH)
        base.font.bold = True
        child = self.doc.styles.add_style('Second toggle', WD_STYLE_TYPE.PARAGRAPH)
        child.base_style = base; child.font.bold = True
        self.doc.add_paragraph('Two toggles', child)
        self.assertFalse(self.inspect().build_settings().bold)
        self.doc = Document()
        self.doc.styles['Normal'].font.bold = True
        p = self.doc.add_paragraph('Override')
        p.runs[0].bold = False
        self.assertFalse(self.inspect().build_settings().bold)

    def test_security_validation_and_changed_source_are_rejected(self):
        from app.profile_import import inspect_profile
        from zipfile import ZipFile
        self.doc.add_paragraph('Original'); self.doc.save(self.source)
        with patch('app.profile_import.fingerprint', side_effect=['before', 'after']):
            with self.assertRaisesRegex(FormatterError, 'mudou'):
                inspect_profile(self.source)
        with ZipFile(self.source, 'a') as archive:
            archive.writestr('word/vbaProject.bin', b'macro')
        with self.assertRaisesRegex(FormatterError, 'macros'):
            inspect_profile(self.source)

    def test_protected_content_ignored_and_table_limitation_disclosed(self):
        p = self.doc.add_paragraph('Corpo')
        self.doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Tabela'
        revision = OxmlElement('w:ins')
        protected = self.doc.add_paragraph('Revisado')
        protected._p.append(revision)
        report = self.inspect()
        self.assertEqual(report.category_options['body'][0].count, 1)
        self.assertIn('table', report.category_options)
        self.assertTrue(any('controladas' in n for n in report.notices))
        self.assertTrue(any('estilos condicionais de tabela' in n for n in report.notices))

    def test_theme_shade_uses_hsl_luminance(self):
        p = self.doc.add_paragraph('Cor sombreada')
        color = OxmlElement('w:color')
        color.set(qn('w:themeColor'), 'accent2')
        color.set(qn('w:themeShade'), 'BF')
        p.runs[0]._r.get_or_add_rPr().append(color)
        report = self.inspect()
        actual = report.build_settings().font_color
        # Microsoft's Word example is 943634; conversion precision can differ
        # by up to two channel units, which is disclosed in the review.
        for i in (0, 2, 4):
            self.assertLessEqual(abs(int(actual[i:i + 2], 16) - int('943634'[i:i + 2], 16)), 2)
        self.assertTrue(any('arredondamento' in n for n in report.notices))

    def test_missing_section_is_rejected_before_review(self):
        self.doc.add_paragraph('Documento sem página')
        section = self.doc.sections[0]._sectPr
        section.getparent().remove(section)
        with self.assertRaisesRegex(FormatterError, 'seção'):
            self.inspect()

    def test_theme_color_honors_document_color_mapping(self):
        mapping = self.doc.settings.element.find(qn('w:clrSchemeMapping'))
        if mapping is None:
            mapping = OxmlElement('w:clrSchemeMapping')
            self.doc.settings.element.append(mapping)
        mapping.set(qn('w:accent1'), 'accent2')
        mapping.set(qn('w:t2'), 'accent1')
        for slot, expected in (('accent1', 'C0504D'), ('text2', '4F81BD')):
            with self.subTest(slot=slot):
                p = self.doc.add_paragraph(slot)
                color = OxmlElement('w:color')
                color.set(qn('w:val'), '000000')
                color.set(qn('w:themeColor'), slot)
                p.runs[0]._r.get_or_add_rPr().append(color)
                actual = self.inspect().build_settings().font_color
                p._p.getparent().remove(p._p)
                self.assertEqual(actual, expected)

    def test_unresolved_theme_keeps_cached_color_without_second_tint(self):
        for rel_id, rel in list(self.doc.part.rels.items()):
            if rel.reltype.endswith('/theme'):
                self.doc.part.drop_rel(rel_id)
        p = self.doc.add_paragraph('Cor armazenada')
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '548DD4')
        color.set(qn('w:themeColor'), 'text2')
        color.set(qn('w:themeTint'), '99')
        p.runs[0]._r.get_or_add_rPr().append(color)
        report = self.inspect()
        self.assertEqual(report.build_settings().font_color, '548DD4')
        self.assertTrue(any('Cor de tema não resolvida' in n for n in report.notices))


if __name__ == '__main__':
    unittest.main()
