import tempfile
import unittest
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from app.config import FormatSettings
from app.document_model import DocumentOverrides, ImageEdit, inspect_document
from app.formatter import format_document, FormatterError


class CategoryFormattingTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / 'input.docx'
        self.doc = Document()
        self.doc.add_heading('Título', 1)
        self.doc.add_paragraph('Corpo')
        self.doc.add_paragraph('Citação', 'Quote')

    def apply(self, settings, overrides=None):
        self.doc.save(self.source)
        result = format_document(self.source, self.root / 'out', settings, overrides=overrides)
        return Document(result.output_path), result

    def test_distinct_rules_and_preservation(self):
        settings = FormatSettings(formatting_mode='by_category', category_rules={
            'title': {'mode':'custom','values':{'font_size':18,'alignment':'center','first_line_indent_cm':0}},
            'quote': {'mode':'preserve'}})
        original = self.doc.paragraphs[2]._p.xml
        out, _ = self.apply(settings)
        self.assertEqual(out.paragraphs[0].runs[0].font.size.pt, 18)
        self.assertEqual(out.paragraphs[1].runs[0].font.size.pt, 12)
        self.assertEqual(out.paragraphs[2]._p.xml, original)

    def test_protected_revision_and_equation_preserved(self):
        ins = OxmlElement('w:ins')
        r = OxmlElement('w:r')
        t = OxmlElement('w:t'); t.text = 'Inserção'
        r.append(t); ins.append(r)
        self.doc.paragraphs[0]._p.append(ins)
        math = OxmlElement('m:oMath')
        self.doc.paragraphs[2]._p.append(math)
        original = self.doc.paragraphs[0]._p.xml
        out, result = self.apply(FormatSettings())
        self.assertEqual(out.paragraphs[0]._p.xml, original)
        self.assertEqual(out.paragraphs[1].runs[0].font.size.pt,12)
        self.assertEqual(len(out._element.body.xpath('.//m:oMath')),1)
        self.assertTrue(any('controladas' in w for w in result.warnings))

    def test_manual_category_and_image_move_with_ratio(self):
        image = self.root / 'test.png'
        Image.new('RGB',(200,100)).save(image)
        self.doc.add_picture(str(image))
        self.doc.save(self.source)
        settings = FormatSettings(formatting_mode='by_category', category_rules={
            'signature': {'mode':'custom','values':{'alignment':'center','font_size':14}}})
        info = inspect_document(self.source, settings)
        edits = DocumentOverrides(info.fingerprint, {1:'signature'}, {0:ImageEdit(4,'center',0,'before')})
        result = format_document(self.source, self.root / 'out', settings, overrides=edits)
        out = Document(result.output_path)
        self.assertEqual(out.paragraphs[0].text,'')
        self.assertAlmostEqual(out.inline_shapes[0].width.cm,4)
        self.assertAlmostEqual(out.inline_shapes[0].height.cm,2)
        self.assertEqual(out.paragraphs[2].runs[0].font.size.pt,14)

    def test_stale_review_rejected(self):
        self.doc.save(self.source)
        info = inspect_document(self.source, FormatSettings())
        self.doc.add_paragraph('Novo'); self.doc.save(self.source)
        with self.assertRaisesRegex(FormatterError,'mudou'):
            format_document(self.source,self.root/'out',FormatSettings(),DocumentOverrides(info.fingerprint))
        self.assertFalse((self.root/'out').exists())

    def test_table_revision_is_unchanged(self):
        cell = self.doc.add_table(rows=1,cols=1).cell(0,0)
        cell.text = 'Texto alterado'
        cell._tc.getparent().get_or_add_trPr().append(OxmlElement('w:ins'))
        original = cell.paragraphs[0]._p.xml
        out, _ = self.apply(FormatSettings())
        self.assertEqual(out.tables[0].cell(0,0).paragraphs[0]._p.xml, original)

    def test_textbox_rule_does_not_leak_from_outer_paragraph(self):
        box = OxmlElement('w:txbxContent'); p = OxmlElement('w:p')
        r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = 'Caixa'
        r.append(t); p.append(r); box.append(p)
        self.doc.paragraphs[1].runs[0]._r.append(box)
        settings = FormatSettings(formatting_mode='by_category', category_rules={
            'textbox': {'mode':'custom','values':{'font_size':9}}})
        out, _ = self.apply(settings)
        self.assertEqual(out.paragraphs[1].runs[0].font.size.pt,12)
        self.assertEqual(out._element.body.xpath('.//w:txbxContent//w:sz')[0].get(qn('w:val')),'18')

    def test_global_image_width_and_alignment_preserve_adjacent_text(self):
        image = self.root/'image.png'; Image.new('RGB',(200,100)).save(image)
        self.doc.paragraphs[1].add_run().add_picture(str(image))
        out, _ = self.apply(FormatSettings(body_image_width_cm=4,body_image_alignment='right'))
        self.assertEqual(out.paragraphs[1].text,'Corpo')
        self.assertEqual(out.paragraphs[2].text,'')
        self.assertAlmostEqual(out.inline_shapes[0].width.cm,4)
        self.assertAlmostEqual(out.inline_shapes[0].height.cm,2)

    def test_invalid_or_overheight_image_never_published(self):
        image = self.root/'image.png'; Image.new('RGB',(10,100)).save(image)
        self.doc.add_picture(str(image)); self.doc.save(self.source)
        with self.assertRaisesRegex(FormatterError,'cabe'):
            format_document(self.source,self.root/'out',FormatSettings(body_image_width_cm=10))
        self.assertFalse((self.root/'out').exists())

    def test_profile_styles_and_manual_override_take_precedence(self):
        settings=FormatSettings(formatting_mode='by_category',style_categories={'Normal':'signature'},
            category_rules={'signature':{'mode':'custom','values':{'font_size':15}}})
        out,_=self.apply(settings)
        self.assertEqual(out.paragraphs[1].runs[0].font.size.pt,15)

    def test_images_keep_order_and_alignment_when_moving(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        image = self.root/'image.png'; Image.new('RGB',(200,100)).save(image)
        p = self.doc.paragraphs[1]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for label in ('FIRST','SECOND'):
            shape = p.add_run().add_picture(str(image)); shape._inline.docPr.set('descr',label)
        self.doc.save(self.source)
        info=inspect_document(self.source,FormatSettings())
        edits=DocumentOverrides(info.fingerprint,image_edits={
            0:ImageEdit(target_paragraph=0),1:ImageEdit(target_paragraph=0)})
        result=format_document(self.source,self.root/'out',FormatSettings(),edits)
        out=Document(result.output_path)
        self.assertEqual([shape._inline.docPr.get('descr') for shape in out.inline_shapes],['FIRST','SECOND'])
        self.assertEqual(out.paragraphs[1].alignment,WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(out.paragraphs[2].alignment,WD_ALIGN_PARAGRAPH.CENTER)

    def test_resized_image_has_no_text_indents(self):
        image=self.root/'image.png'; Image.new('RGB',(200,100)).save(image)
        self.doc.add_picture(str(image))
        out,_=self.apply(FormatSettings(body_image_width_cm=16))
        from docx.text.paragraph import Paragraph
        from app.document_model import nearest_paragraph
        p=Paragraph(nearest_paragraph(out.inline_shapes[0]._inline),out._body)
        self.assertEqual(p.paragraph_format.first_line_indent.cm,0)
        self.assertEqual(p.paragraph_format.left_indent.cm,0)
        self.assertEqual(p.paragraph_format.right_indent.cm,0)

    def test_strict_mode_rejects_cell_revision(self):
        cell=self.doc.add_table(rows=1,cols=1).cell(0,0)
        cell._tc.get_or_add_tcPr().append(OxmlElement('w:cellIns'))
        self.doc.save(self.source)
        with self.assertRaisesRegex(FormatterError,'controladas'):
            format_document(self.source,self.root/'out',FormatSettings(complex_content_mode='reject'))
