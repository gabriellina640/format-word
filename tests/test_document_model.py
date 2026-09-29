import tempfile
import unittest
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from PIL import Image
from app.config import FormatSettings
from app.document_model import inspect_document, classify_paragraph, own_runs, is_protected


class DocumentModelTests(unittest.TestCase):
    def test_styles_tables_and_image_inspection(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image = root / 'image.png'
            Image.new('RGB', (200, 100)).save(image)
            doc = Document()
            doc.add_heading('Título', 1)
            doc.add_paragraph('Texto')
            doc.add_paragraph('Citação', 'Quote')
            doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Tabela'
            doc.add_picture(str(image))
            path = root / 'test.docx'
            doc.save(path)
            info = inspect_document(path, FormatSettings())
            self.assertEqual([p.category for p in info.paragraphs[:4]], ['title', 'body', 'quote', 'table'])
            self.assertEqual(len(info.fingerprint), 64)
            self.assertEqual(len(info.images), 1)
            self.assertAlmostEqual(info.images[0].width_cm / info.images[0].height_cm, 2)
            self.assertTrue(info.images[0].editable)

    def test_revision_and_nested_textbox_runs(self):
        doc = Document()
        p = doc.add_paragraph('Fora')
        ins = OxmlElement('w:ins')
        ins.append(OxmlElement('w:r'))
        p._p.append(ins)
        self.assertTrue(is_protected(p._p))
        outer = doc.add_paragraph('Externo')
        box = OxmlElement('w:txbxContent')
        inner = OxmlElement('w:p')
        inner.append(OxmlElement('w:r'))
        box.append(inner)
        outer.runs[0]._r.append(box)
        self.assertEqual(len(own_runs(outer._p)), 1)
        self.assertEqual(len(own_runs(inner)), 1)
        self.assertEqual(classify_paragraph(inner, doc, FormatSettings()), 'textbox')

    def test_table_revision_properties_protect_paragraphs(self):
        for name, container in [('ins','tr'),('del','tr'),('trPrChange','tr'),('tcPrChange','tc'),('cellIns','tc')]:
            doc = Document(); cell = doc.add_table(rows=1, cols=1).cell(0,0)
            if container == 'tr':
                properties = cell._tc.getparent().get_or_add_trPr()
            else:
                properties = cell._tc.get_or_add_tcPr()
            properties.append(OxmlElement('w:' + name))
            with self.subTest(name=name):
                self.assertTrue(is_protected(cell.paragraphs[0]._p))

    def test_custom_and_inherited_style_classification(self):
        from docx.enum.style import WD_STYLE_TYPE
        doc = Document()
        custom = doc.styles.add_style('Título local', WD_STYLE_TYPE.PARAGRAPH)
        custom.base_style = doc.styles['Heading 2']
        p = doc.add_paragraph('Texto', custom)
        self.assertEqual(classify_paragraph(p._p, doc, FormatSettings()), 'title')
        settings = FormatSettings(style_categories={custom.style_id:'signature'})
        self.assertEqual(classify_paragraph(p._p, doc, settings), 'signature')
        settings.style_categories = {custom.name:'quote'}
        self.assertEqual(classify_paragraph(p._p, doc, settings), 'quote')

    def test_drawing_with_nested_picture_is_not_counted_as_picture(self):
        from app.document_model import image_elements
        doc = Document(); p = doc.add_paragraph()
        outer = OxmlElement('wp:inline')
        box = OxmlElement('w:txbxContent'); inner_p = OxmlElement('w:p')
        inner = OxmlElement('wp:inline')
        graphic = OxmlElement('a:graphic'); data = OxmlElement('a:graphicData')
        data.append(OxmlElement('pic:pic')); graphic.append(data); inner.append(graphic)
        inner_p.append(inner); box.append(inner_p); outer.append(box); p._p.append(outer)
        self.assertEqual(image_elements(doc), [inner])
